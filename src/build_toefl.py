#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""../新东方托福【绿宝书】-托福词汇.xlsx -> data/toefl.js（第十二次优化 2）

原表一张 Sheet 一个 List，48 个 List，逐张表对应生成一个单元（u 1001-1048）。
i 从 5001 起、u 从 1001 起 —— 都是刻意跳开的：现有 SAT 数据用到 i≤2538 / u≤37，
这里没有贴着接上去，是想给「以后又给 SAT 那边加错词本一类的东西」留出空间；
往后要是真加了，只要新数据的 i 不越过 5000、u 不越过 1000，这两批就不会撞上。

校验规则照抄 app.js 里「导入词表」用的那一套（parseTable / wordCell）：
单词要以字母开头、不含中文、最多 6 个词；释义必须含中文字符，没有就不是真释义
（第十二次优化 2 就是栽在这条上：原 Python 脚本没查这条，把两处空白的格式残留、
两处错填的英文例句当成释义直接收了进去，后来对着 app.js 自己的解析结果一比对才发现）。
不满足就跳过并打印警告——不静默吞掉，免得下次改了源表又recreate 出错数据。

WORD_FIXES / MEANING_FIXES 是这一轮手动订正过的已知坏格：
· sheet 03 "confl ict"：原表词形中间多了个空格，是抄录问题，不是词组。
· sheet 32/41 acid/drain：意思格是空的 HTML 格式残留。
· sheet 40/43 conceivable/delirium：意思格错填成了英文例句。
这四条的音标都是原表真实数据，照单收下；释义是查漏之后按词义手动补的，不是原书内容。
"""
import os
import re
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(os.path.dirname(ROOT), "新东方托福【绿宝书】-托福词汇.xlsx")
OUT = os.path.join(ROOT, "data", "toefl.js")

I_START = 5001
U_START = 1001

CJK = re.compile(r"[一-鿿]")


def phrase_ok(w):
    return bool(w) and re.match(r"^[A-Za-z]", w) and len(w) <= 60 \
        and not CJK.search(w) and len(w.split()) <= 6


WORD_FIXES = {
    ("03", "confl ict"): "conflict",
}
MEANING_FIXES = {
    ("32", "acid"): "adj. 酸的，尖刻的(sour, sharp) n. 酸，酸性物质",
    ("40", "conceivable"): "adj. 可想象的，可能的(imaginable, possible)",
    ("41", "drain"): "v. 排干，耗尽(deplete, exhaust) n. 排水沟，下水道",
    ("43", "delirium"): "n. 谵妄，精神错乱(frenzy)；狂喜，极度兴奋",
}

wb = openpyxl.load_workbook(SRC, data_only=True)

words = []
units = []
i = I_START
u = U_START
warnings = []

for si, sn in enumerate(wb.sheetnames):
    ws = wb[sn]
    unit_words = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        word = r[1]
        if word is None or not str(word).strip():
            continue
        w = str(word).strip()
        w = WORD_FIXES.get((sn, w), w)
        if not phrase_ok(w):
            warnings.append(f"sheet {sn}: 跳过「{w}」—— 不像一个单词/词组")
            continue
        us_ipa = str(r[2]).strip() if r[2] else ""
        uk_ipa = str(r[3]).strip() if r[3] else ""
        p = (us_ipa or uk_ipa).strip("/[] ")
        z = str(r[4]).strip() if r[4] else ""
        z = MEANING_FIXES.get((sn, w), z)
        if not CJK.search(z):
            warnings.append(f"sheet {sn}: 跳过「{w}」—— 释义格里没有中文：{z!r}")
            continue
        rec = {"i": i, "u": u, "w": w, "p": p, "z": z}
        words.append(rec)
        unit_words.append(rec)
        i += 1
    if not unit_words:
        continue
    units.append({
        "u": u, "n": len(unit_words),
        "a": unit_words[0]["w"], "b": unit_words[-1]["w"],
        "t": "托福绿宝书 · List %02d" % (si + 1),
    })
    u += 1

import json
dump = lambda o: json.dumps(o, ensure_ascii=False, separators=(",", ":"))
with open(OUT, "w", encoding="utf-8") as f:
    f.write("/* 新东方《绿宝书》托福词汇：48 个 List，逐张表对应一个单元（第十二次优化 2）。\n")
    f.write("   由 src/build_toefl.py 从 ../新东方托福【绿宝书】-托福词汇.xlsx 生成，\n")
    f.write("   i 从 %d 起、u 从 %d 起，跟原有 SAT 数据（i 1-2538 · u 1-37）完全分开。 */\n" % (I_START, U_START))
    f.write("window.SAT_WORDS=(window.SAT_WORDS||[]).concat([\n")
    for r in words:
        f.write(dump(r) + ",\n")
    f.write("]);\n")
    f.write("window.SAT_UNITS=(window.SAT_UNITS||[]).concat(" + dump(units) + ");\n")

print(f"{len(words)} words -> {len(units)} units (i {words[0]['i']}-{words[-1]['i']}, u {units[0]['u']}-{units[-1]['u']})")
if warnings:
    print(f"\n{len(warnings)} 条警告（没有静默跳过，都在这）：")
    for wmsg in warnings:
        print("  " + wmsg)
